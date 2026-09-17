"""Importa el maestro matrículas × combustible (xlsx) a matriculas_combustible.

Uso:
  python -m scripts.import_maestro_matriculas
  python -m scripts.import_maestro_matriculas --replace
  python -m scripts.import_maestro_matriculas --path /ruta/al.xlsx --replace

Por defecto: data/maestro-aviones-version-final.xlsx
  (fallback: data/maestro-matriculas-combustible.xlsx)

Formato FINAL (sheet BASE / BASE FINAL):
  CodigoProducto | Combustible | Matricula | Avion

Formato legacy (sheet MATRICULAS Y COMBUSTIBLE) aún soportado.

--replace: FULL REPLACE del maestro (borra filas previas y carga el archivo).
"""

from __future__ import annotations

import argparse
import logging
import os
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

DEFAULT_PATH = ROOT / "data" / "maestro-aviones-version-final.xlsx"

logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
log = logging.getLogger("import_maestro")


@dataclass
class Candidate:
    key: str
    display: str
    grado: str
    modelo: str


def _pick_sheet(wb) -> str:
    for name in ("BASE FINAL", "BASE", "MATRICULAS Y COMBUSTIBLE"):
        if name in wb.sheetnames:
            return name
    return wb.sheetnames[0]


def read_candidates(path: Path) -> tuple[list[Candidate], dict]:
    from openpyxl import load_workbook

    from app.matricula import grado_from_producto_nombre, normalize_matricula

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = _pick_sheet(wb)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    header = [str(c or "").strip() for c in next(rows)]
    idx = {h: i for i, h in enumerate(header)}
    stats = Counter()
    stats["sheet"] = sheet
    by_key: dict[str, Candidate] = {}

    def col(*names, default=None):
        for n in names:
            if n in idx:
                return idx[n]
        return default

    i_mat = col("Matricula", "Matricula_Validada")
    i_mat_fb = col("Matricula")
    i_mat_val = col("Matricula_Validada")
    i_comb = col("Combustible", "ProductoNombre")
    i_avion = col("Avion", "Modelo_Avion")

    for raw in rows:
        stats["rows_raw"] += 1
        if not raw:
            continue
        # Prefer Matricula_Validada if present in legacy
        mat_raw = None
        if i_mat_val is not None and i_mat_val < len(raw) and raw[i_mat_val] not in (None, ""):
            mat_raw = raw[i_mat_val]
        elif i_mat is not None and i_mat < len(raw):
            mat_raw = raw[i_mat]
        elif i_mat_fb is not None and i_mat_fb < len(raw):
            mat_raw = raw[i_mat_fb]

        comb_raw = raw[i_comb] if i_comb is not None and i_comb < len(raw) else ""
        modelo = ""
        if i_avion is not None and i_avion < len(raw) and raw[i_avion] not in (None, ""):
            modelo = str(raw[i_avion]).strip()
            if modelo.startswith("#REF"):
                modelo = ""
            modelo = modelo[:80]

        key = normalize_matricula(str(mat_raw or ""))
        if not key:
            stats["skip_empty_matricula"] += 1
            continue

        grado = grado_from_producto_nombre(str(comb_raw or ""))
        if not grado:
            # Combustible may already be "JET A-1" / "AVGAS 100LL"
            cs = str(comb_raw or "").strip().upper()
            if "JET" in cs:
                grado = "JET A-1"
            elif "AVGAS" in cs or "100LL" in cs:
                grado = "AVGAS 100LL"
        if not grado:
            stats["skip_unknown_grado"] += 1
            continue

        display = str(mat_raw or "").strip().upper() or key
        if key in by_key:
            stats["duplicate_matriculas"] += 1
            continue
        by_key[key] = Candidate(key=key, display=display[:40], grado=grado, modelo=modelo)
        stats["rows_valid"] += 1
        if modelo:
            stats["with_avion"] += 1
        if grado == "JET A-1":
            stats["jet"] += 1
        elif grado == "AVGAS 100LL":
            stats["avgas"] += 1

    wb.close()
    chosen = list(by_key.values())
    stats["unique_matriculas"] = len(chosen)
    return chosen, dict(stats)


def import_candidates(
    candidates: list[Candidate], *, dry_run: bool = False, replace: bool = False
) -> dict:
    from sqlalchemy import delete, func, select

    from app.database import SessionLocal
    from app.matricula import upsert_matricula_combustible
    from app.models import MatriculaCombustible

    inserted = updated = deleted = 0
    db = SessionLocal()
    try:
        if replace and not dry_run:
            deleted = db.scalar(select(func.count()).select_from(MatriculaCombustible)) or 0
            db.execute(delete(MatriculaCombustible))
            db.flush()
        existing = {
            row.matricula: row
            for row in db.scalars(select(MatriculaCombustible)).all()
        }
        for cand in candidates:
            row = existing.get(cand.key)
            if row is None:
                inserted += 1
                if not dry_run:
                    upsert_matricula_combustible(
                        db,
                        raw_matricula=cand.display,
                        combustible=cand.grado,
                        modelo=cand.modelo,
                        tipo=cand.modelo or None,
                        activo=True,
                    )
            else:
                changed = (
                    (row.combustible or "") != cand.grado
                    or (row.modelo or "") != (cand.modelo or "")
                    or (row.tipo or "") != (cand.modelo or "")
                    or not row.activo
                    or (row.matricula_display or "") != cand.display
                )
                if changed:
                    updated += 1
                    if not dry_run:
                        upsert_matricula_combustible(
                            db,
                            raw_matricula=cand.display,
                            combustible=cand.grado,
                            modelo=cand.modelo or None,
                            tipo=cand.modelo or None,
                            activo=True,
                        )
        if not dry_run:
            db.commit()
        total = db.scalar(select(func.count()).select_from(MatriculaCombustible))
    finally:
        db.close()

    return {
        "inserted": inserted,
        "updated": updated,
        "deleted_before_load": deleted if replace else 0,
        "unchanged": len(candidates) - inserted - updated,
        "total_in_db": total if not dry_run else None,
        "dry_run": dry_run,
        "replace": replace,
    }


def resolve_path(cli_path: str | None) -> Path:
    raw = cli_path or os.environ.get("MAESTRO_MATRICULAS_PATH") or str(DEFAULT_PATH)
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = (Path.cwd() / path).resolve()
    if not path.exists() and not cli_path and not os.environ.get("MAESTRO_MATRICULAS_PATH"):
        alt = ROOT / "data" / "maestro-matriculas-combustible.xlsx"
        if alt.exists():
            path = alt
    if not path.exists():
        raise SystemExit(f"No existe el archivo: {path}")
    return path


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Import maestro matrículas + combustible")
    parser.add_argument("--path", help="Ruta al xlsx")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument(
        "--replace",
        action="store_true",
        help="Full replace: borra el maestro y carga el archivo",
    )
    args = parser.parse_args(argv)

    path = resolve_path(args.path)
    log.info("Leyendo %s", path)
    candidates, read_stats = read_candidates(path)
    log.info("Lectura: %s", read_stats)
    result = import_candidates(candidates, dry_run=args.dry_run, replace=args.replace)
    log.info("Import: %s", result)
    print(
        f"OK unique={read_stats.get('unique_matriculas')} "
        f"JET={read_stats.get('jet')} AVGAS={read_stats.get('avgas')} "
        f"inserted={result['inserted']} updated={result['updated']} "
        f"deleted={result['deleted_before_load']} total_db={result['total_in_db']} "
        f"replace={args.replace} dry_run={args.dry_run}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
